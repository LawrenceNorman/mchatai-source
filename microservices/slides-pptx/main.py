"""Slides PPTX — AI Slides → editable PowerPoint, on the user's Mac (Phase SL2.9).

Mounts at `/svc/slides-pptx/*` in the mChatAIShell sidecar. The app sends the deck already
interpreted (theme + brand resolved, charts resolved to numbers); this module only draws it.
Positions mirror AISlides/SlideRenderView.swift in its 1920×1080 design space.
"""

from __future__ import annotations

import importlib
import os
import re
import subprocess
import sys
from pathlib import Path
from typing import Any, Optional

from fastapi import APIRouter, HTTPException
from pydantic import BaseModel

router = APIRouter()

DESIGN_W, DESIGN_H = 1920.0, 1080.0
SLIDE_W_IN, SLIDE_H_IN = 13.333, 7.5
MARGIN = 110.0

FONT_FAMILIES = {
    "default": "Helvetica Neue",
    "rounded": "Arial Rounded MT Bold",
    "serif": "Georgia",
    "monospaced": "Menlo",
}


class ExportRequest(BaseModel):
    deck: dict
    output_path: str
    assets_dir: Optional[str] = None


# --------------------------------------------------------------------------- helpers


def _pptx():
    """python-pptx, installing it on first use when the runtime does not have it.

    mChatAIShell installs every service's python_deps in ONE pip call that it treats as
    non-fatal and cuts off at 120 s; on a machine with many services that call can die before
    it reaches python-pptx, and the service then mounts without its one dependency. Rather than
    fail every export until someone repairs the runtime by hand, install our own dependency
    once, into the interpreter we are running in, and say so clearly if that fails too."""
    try:
        import pptx  # noqa: F401
        return pptx
    except ImportError:
        pass
    try:
        result = subprocess.run([sys.executable, "-m", "pip", "install", "--quiet", "python-pptx>=1.0"],
                                capture_output=True, text=True, timeout=110)
    except subprocess.TimeoutExpired as exc:
        raise HTTPException(status_code=412, detail="installing python-pptx timed out — try the export again") from exc
    if result.returncode != 0:
        raise HTTPException(status_code=412,
                            detail="python-pptx could not be installed: " + (result.stderr or result.stdout)[-400:])
    importlib.invalidate_caches()
    import pptx  # noqa: F401
    return pptx


def _emu(design: float) -> int:
    return int(design / DESIGN_W * SLIDE_W_IN * 914400)


def _pt(design_px: float):
    from pptx.util import Pt
    return Pt(design_px * SLIDE_W_IN * 72.0 / DESIGN_W)


def _rgb(hex_value: Optional[str], fallback: str = "#111214"):
    from pptx.dml.color import RGBColor
    raw = (hex_value or fallback).lstrip("#")
    if len(raw) != 6:
        raw = fallback.lstrip("#")
    return RGBColor.from_string(raw.upper())


INLINE = re.compile(r"\*\*\*(.+?)\*\*\*|\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`|\[(.+?)\]\((\S+?)\)")


def _runs(text: str):
    """Inline Markdown → (text, bold, italic, code, link) runs. Unmatched markers stay literal."""
    pos = 0
    for match in INLINE.finditer(text or ""):
        if match.start() > pos:
            yield text[pos:match.start()], False, False, False, None
        both, strong, emph, code, link_text, link_url = match.groups()
        if both is not None:
            yield both, True, True, False, None
        elif strong is not None:
            yield strong, True, False, False, None
        elif emph is not None:
            yield emph, False, True, False, None
        elif code is not None:
            yield code, False, False, True, None
        else:
            yield link_text, False, False, False, link_url
        pos = match.end()
    if pos < len(text or ""):
        yield text[pos:], False, False, False, None


class Ctx:
    def __init__(self, deck: dict, assets_dir: Optional[str]):
        theme = deck.get("theme") or {}
        self.theme = theme
        self.brand = deck.get("brand") or {}
        self.assets = Path(assets_dir) if assets_dir else None
        self.fg = theme.get("text", "#111214")
        self.muted = theme.get("muted", "#6b7280")
        self.accent = theme.get("accent", "#4f6df5")
        self.accent_contrast = theme.get("accentContrast", "#ffffff")
        self.surface = theme.get("surface", "#f2f2f4")
        self.background = theme.get("background", "#ffffff")
        self.display = FONT_FAMILIES.get(theme.get("displayFontDesign", "default"), "Helvetica Neue")
        self.body = FONT_FAMILIES.get(theme.get("bodyFontDesign", "default"), "Helvetica Neue")
        self.palette = theme.get("palette") or ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300"]
        self.warnings: list[str] = []

    def picture(self, name: Optional[str]) -> Optional[Path]:
        if not name or not self.assets:
            return None
        path = self.assets / name
        return path if path.exists() else None


def _fill_text(frame, text: str, ctx: Ctx, size: float, color: str, *, bold=False, italic=False,
               font: Optional[str] = None, align=None):
    from pptx.enum.text import PP_ALIGN
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    # Always explicit: a title placeholder otherwise inherits the template's centring, and the
    # app left-aligns titles on every layout but the cover and the big fact.
    paragraph.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(align or "", PP_ALIGN.LEFT)
    for chunk, r_bold, r_italic, r_code, link in _runs(text):
        run = paragraph.add_run()
        run.text = chunk
        run.font.size = _pt(size)
        run.font.bold = bold or r_bold
        run.font.italic = italic or r_italic
        run.font.name = "Menlo" if r_code else (font or ctx.body)
        run.font.color.rgb = _rgb(ctx.accent if link else color)
        if link:
            run.hyperlink.address = link
            run.font.underline = True


def _text_box(slide, ctx: Ctx, x, y, w, h, text: str, size: float, color: str, **kwargs):
    box = slide.shapes.add_textbox(_emu(x), _emu(y), _emu(w), _emu(h))
    box.text_frame.margin_left = box.text_frame.margin_right = 0
    box.text_frame.margin_top = box.text_frame.margin_bottom = 0
    _fill_text(box.text_frame, text, ctx, size, color, **kwargs)
    return box


def _bullets(slide, ctx: Ctx, x, y, w, h, lines: list[str], size: float = 38, color: Optional[str] = None,
             heading: Optional[str] = None):
    """A bullet list as real paragraphs (editable, one bullet each) with the accent dot."""
    from pptx.oxml.ns import qn
    from lxml import etree
    box = slide.shapes.add_textbox(_emu(x), _emu(y), _emu(w), _emu(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    first = True
    items = ([("heading", heading)] if heading else []) + [("bullet", line) for line in lines[:8]]
    for kind, line in items:
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        for chunk, r_bold, r_italic, r_code, link in _runs(line):
            run = paragraph.add_run()
            run.text = chunk
            is_heading = kind == "heading"
            run.font.size = _pt(40 if is_heading else size)
            run.font.bold = is_heading or r_bold
            run.font.italic = r_italic
            run.font.name = "Menlo" if r_code else ctx.body
            run.font.color.rgb = _rgb(ctx.accent if (is_heading or link) else (color or ctx.fg))
            if link:
                run.hyperlink.address = link
        paragraph.space_after = _pt(26 if kind == "heading" else 30)
        if kind == "bullet":
            p_pr = paragraph._p.get_or_add_pPr()
            p_pr.set("marL", str(_emu(37)))
            p_pr.set("indent", str(-_emu(37)))
            bu_clr = etree.SubElement(p_pr, qn("a:buClr"))
            srgb = etree.SubElement(bu_clr, qn("a:srgbClr"))
            srgb.set("val", ctx.accent.lstrip("#").upper())
            bu_font = etree.SubElement(p_pr, qn("a:buFont"))
            bu_font.set("typeface", "Arial")
            bu_char = etree.SubElement(p_pr, qn("a:buChar"))
            bu_char.set("char", "•")
    return box


def _rect(slide, x, y, w, h, color: str, *, radius: bool = False, alpha: Optional[int] = None):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn
    from lxml import etree
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   _emu(x), _emu(y), _emu(w), _emu(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    if alpha is not None:
        srgb = shape.fill._xPr.find(".//" + qn("a:srgbClr"))
        if srgb is not None:
            node = etree.SubElement(srgb, qn("a:alpha"))
            node.set("val", str(alpha * 1000))
    return shape


def _oval(slide, x, y, d, color: str):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, _emu(x), _emu(y), _emu(d), _emu(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _picture(slide, ctx: Ctx, ref: Optional[dict], x, y, w, h):
    """A picture filling (x,y,w,h) — cropped to its focus — or, for "fit", letterboxed whole."""
    path = ctx.picture((ref or {}).get("file"))
    if path is None:
        _rect(slide, x, y, w, h, ctx.surface, radius=True)
        prompt = (ref or {}).get("prompt")
        if prompt:
            _text_box(slide, ctx, x + 40, y + h / 2 - 40, w - 80, 80, prompt, 26, ctx.muted, align="center")
        ctx.warnings.append("a picture was missing — drew a placeholder")
        return
    from PIL import Image  # python-pptx depends on Pillow
    with Image.open(path) as im:
        iw, ih = im.size
    if (ref or {}).get("fit") == "fit":
        scale = min(w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        _rect(slide, x, y, w, h, ctx.surface)
        slide.shapes.add_picture(str(path), _emu(x + (w - pw) / 2), _emu(y + (h - ph) / 2), _emu(pw), _emu(ph))
        return
    focus = (ref or {}).get("focus") or {}
    fx, fy = float(focus.get("x", 0.5)), float(focus.get("y", 0.5))
    pic = slide.shapes.add_picture(str(path), _emu(x), _emu(y), _emu(w), _emu(h))
    frame_ratio, image_ratio = w / h, iw / ih
    if image_ratio > frame_ratio:            # too wide: crop left/right around the focus
        keep = frame_ratio / image_ratio
        spare = 1 - keep
        pic.crop_left = spare * fx
        pic.crop_right = spare * (1 - fx)
    elif image_ratio < frame_ratio:          # too tall: crop top/bottom
        keep = image_ratio / frame_ratio
        spare = 1 - keep
        pic.crop_top = spare * fy
        pic.crop_bottom = spare * (1 - fy)


def _title(slide, ctx: Ctx, text: Optional[str], *, size=60, x=MARGIN, y=MARGIN - 10, w=DESIGN_W - 2 * MARGIN,
           h=150, color=None, align=None):
    """Uses the layout's REAL title placeholder when there is one, so PowerPoint's outline
    view and accessibility tools see the slide title."""
    if not text:
        return
    placeholder = slide.shapes.title
    if placeholder is not None:
        placeholder.left, placeholder.top = _emu(x), _emu(y)
        placeholder.width, placeholder.height = _emu(w), _emu(h)
        frame = placeholder.text_frame
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        from pptx.enum.text import MSO_ANCHOR
        frame.vertical_anchor = MSO_ANCHOR.TOP
        _fill_text(frame, text, ctx, size, color or ctx.fg, bold=True, font=ctx.display, align=align)
    else:
        _text_box(slide, ctx, x, y, w, h, text, size, color or ctx.fg, bold=True, font=ctx.display, align=align)


# --------------------------------------------------------------------------- charts


def _plain_frame(chart):
    """No fill and no outline on the chart area or the plot area — the slide's own rounded
    surface is the frame. python-pptx has no API for either, so the c:spPr is inserted at the
    schema's position (after c:chart; before c:extLst in the plot area)."""
    from lxml import etree
    from pptx.oxml.ns import qn

    def no_fill():
        sp_pr = etree.Element(qn("c:spPr"))
        etree.SubElement(sp_pr, qn("a:noFill"))
        line = etree.SubElement(sp_pr, qn("a:ln"))
        etree.SubElement(line, qn("a:noFill"))
        return sp_pr

    space = chart._chartSpace
    chart_el = space.find(qn("c:chart"))
    for old in space.findall(qn("c:spPr")):
        space.remove(old)
    chart_el.addnext(no_fill())
    plot_area = chart_el.find(qn("c:plotArea"))
    if plot_area is not None:
        for old in plot_area.findall(qn("c:spPr")):
            plot_area.remove(old)
        ext = plot_area.find(qn("c:extLst"))
        if ext is not None:
            ext.addprevious(no_fill())
        else:
            plot_area.append(no_fill())


def _chart(slide, ctx: Ctx, chart: dict, x, y, w, h):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

    _rect(slide, x, y, w, h, ctx.surface, radius=True)
    x, y, w, h = x + 40, y + 40, w - 80, h - 80
    kind = chart.get("type", "bar")

    if chart.get("empty"):
        _text_box(slide, ctx, x, y + h / 2 - 60, w, 120, chart["empty"], 30, ctx.muted, align="center")
        return
    if kind == "stat" and chart.get("stat"):
        stat = chart["stat"]
        _text_box(slide, ctx, x, y + 40, w, 200, stat.get("big", ""), 150, ctx.accent, bold=True,
                  font=ctx.display, align="center")
        _text_box(slide, ctx, x, y + 250, w, 60, stat.get("label", ""), 34, ctx.fg, align="center")
        tiles = stat.get("tiles", [])
        tile_w = w / max(len(tiles), 1)
        for i, (label, value) in enumerate(tiles):
            _text_box(slide, ctx, x + i * tile_w, y + 360, tile_w, 60, value, 40, ctx.fg, bold=True, align="center")
            _text_box(slide, ctx, x + i * tile_w, y + 420, tile_w, 40, label, 24, ctx.muted, align="center")
        return
    if kind == "table" and chart.get("table"):
        table = chart["table"]
        columns, rows = table.get("columns", []), table.get("rows", [])
        numeric = table.get("numeric", [False] * len(columns))
        shape = slide.shapes.add_table(len(rows) + 1, len(columns), _emu(x), _emu(y), _emu(w),
                                       _emu(min(h, 70 * (len(rows) + 1))))
        grid = shape.table
        from pptx.enum.text import PP_ALIGN
        for c, name in enumerate(columns):
            cell = grid.cell(0, c)
            _fill_text(cell.text_frame, str(name), ctx, 26, ctx.accent_contrast, bold=True)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(ctx.accent)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row):
                cell = grid.cell(r, c)
                _fill_text(cell.text_frame, str(value), ctx, 24, ctx.fg)
                if c < len(numeric) and numeric[c]:
                    cell.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(ctx.surface if r % 2 else ctx.background)
        return

    categories = [str(c) for c in chart.get("categories", [])]
    values = [float(v) for v in chart.get("values", [])]
    if not categories or not values:
        _text_box(slide, ctx, x, y + h / 2 - 40, w, 80, "No data", 30, ctx.muted, align="center")
        return
    data = CategoryChartData()
    data.categories = categories
    data.add_series(chart.get("seriesName") or "Value", values)
    chart_type = {
        "bar": XL_CHART_TYPE.BAR_CLUSTERED if chart.get("grouped") else XL_CHART_TYPE.COLUMN_CLUSTERED,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "line": XL_CHART_TYPE.LINE,
        "area": XL_CHART_TYPE.AREA,
        "points": XL_CHART_TYPE.LINE_MARKERS,
        "pie": XL_CHART_TYPE.PIE,
        "donut": XL_CHART_TYPE.DOUGHNUT,
    }.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
    if kind in ("pie", "donut"):
        # The slide's look: the ring on the left, a legend of swatch · name · share on the right,
        # drawn as shapes. A chart legend is skipped on purpose — Apple's Office importer (Quick
        # Look, Finder previews) drops the WHOLE chart when a legend sits outside the plot.
        side = min(h, w * 0.5)
        graphic = slide.shapes.add_chart(chart_type, _emu(x), _emu(y + (h - side) / 2), _emu(side), _emu(side), data)
        plot_chart = graphic.chart
        plot_chart.has_legend = False
        plot = plot_chart.plots[0]
        plot.vary_by_categories = True
        colors = []
        points = plot.series[0].points
        for i in range(len(categories)):
            is_other = categories[i] == "Other" and i == len(categories) - 1 and len(categories) > 1
            color = ctx.muted if is_other else ctx.palette[min(i, len(ctx.palette) - 1)]
            colors.append(color)
            fill = points[i].format.fill
            fill.solid()
            fill.fore_color.rgb = _rgb(color)
        _plain_frame(plot_chart)
        total = sum(values) or 1
        row_h = min(64.0, h / max(len(categories), 1))
        legend_x = x + side + 72
        start_y = y + (h - row_h * len(categories)) / 2
        for i, (name, value) in enumerate(zip(categories, values)):
            ly = start_y + i * row_h
            _rect(slide, legend_x, ly + (row_h - 30) / 2, 30, 30, colors[i], radius=True)
            _text_box(slide, ctx, legend_x + 50, ly + (row_h - 40) / 2, 520, 40, name, 28, ctx.fg)
            _text_box(slide, ctx, legend_x + 580, ly + (row_h - 40) / 2, 140, 40,
                      f"{round(value / total * 100)}%", 26, ctx.muted, align="right")
        return
    graphic = slide.shapes.add_chart(chart_type, _emu(x), _emu(y), _emu(w), _emu(h), data)
    plot_chart = graphic.chart
    plot_chart.font.size = _pt(24)
    plot_chart.font.color.rgb = _rgb(ctx.fg)
    plot_chart.font.name = ctx.body
    _plain_frame(plot_chart)
    plot_chart.has_legend = False
    series = plot_chart.plots[0].series[0]
    fmt = series.format
    if kind in ("line", "points"):
        fmt.line.color.rgb = _rgb(ctx.accent)
        fmt.line.width = _pt(5)
    else:
        fmt.fill.solid()
        fmt.fill.fore_color.rgb = _rgb(ctx.accent)
    if kind in ("bar", "column") and chart.get("grouped"):
        plot = plot_chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.number_format = "#,##0.##"
        plot.data_labels.number_format_is_linked = False
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        plot.gap_width = 60
        try:
            plot_chart.value_axis.visible = False
            plot_chart.value_axis.has_major_gridlines = False
        except Exception:  # noqa: BLE001 - axis access differs per chart type
            pass


# --------------------------------------------------------------------------- layouts


def _layout_title(slide, s, ctx):
    _rect(slide, DESIGN_W / 2 - 80, 360, 160, 7, ctx.accent)
    _title(slide, ctx, s.get("title"), size=96, x=MARGIN, y=400, w=DESIGN_W - 2 * MARGIN, h=260, align="center")
    if s.get("subtitle"):
        _text_box(slide, ctx, MARGIN, 690, DESIGN_W - 2 * MARGIN, 120, s["subtitle"], 42, ctx.muted, align="center")


def _layout_section(slide, s, ctx):
    _rect(slide, MARGIN, 330, 14, 420, ctx.accent)
    _title(slide, ctx, s.get("title"), size=84, x=MARGIN + 70, y=380, w=1500, h=220)
    if s.get("subtitle"):
        _text_box(slide, ctx, MARGIN + 70, 620, 1500, 120, s["subtitle"], 40, ctx.muted)


def _layout_bullets(slide, s, ctx):
    _title(slide, ctx, s.get("title"))
    _bullets(slide, ctx, MARGIN, 300, DESIGN_W - 2 * MARGIN, 680, s.get("body") or [])


def _layout_two_column(slide, s, ctx):
    _title(slide, ctx, s.get("title"))
    col_w = (DESIGN_W - 2 * MARGIN - 72) / 2
    for i, key in enumerate(("left", "right")):
        lines = s.get(key) or []
        if lines:
            _bullets(slide, ctx, MARGIN + i * (col_w + 72), 300, col_w, 680, lines[1:], size=34, heading=lines[0])


def _layout_image_side(slide, s, ctx, image_left: bool):
    _title(slide, ctx, s.get("title"))
    image_x = MARGIN if image_left else DESIGN_W - MARGIN - 760
    text_x = MARGIN + 760 + 72 if image_left else MARGIN
    _picture(slide, ctx, s.get("image"), image_x, 300, 760, 620)
    _bullets(slide, ctx, text_x, 300, DESIGN_W - 2 * MARGIN - 760 - 72, 620, s.get("body") or [], size=36)


def _layout_full_bleed(slide, s, ctx):
    _picture(slide, ctx, s.get("image"), 0, 0, DESIGN_W, DESIGN_H)
    if s.get("title") or s.get("caption"):
        _rect(slide, 0, DESIGN_H - 330, DESIGN_W, 330, "#000000", alpha=55)
    if s.get("title"):
        _text_box(slide, ctx, MARGIN * 0.7, DESIGN_H - 300, DESIGN_W - MARGIN * 1.4, 170, s["title"], 68,
                  "#ffffff", bold=True, font=ctx.display)
    if s.get("caption"):
        _text_box(slide, ctx, MARGIN * 0.7, DESIGN_H - 130, DESIGN_W - MARGIN * 1.4, 80, s["caption"], 32, "#e5e5e5")


def _layout_quote(slide, s, ctx):
    _text_box(slide, ctx, MARGIN * 1.3, 110, 300, 260, "“", 220, ctx.accent, bold=True, font=ctx.display)
    _text_box(slide, ctx, MARGIN * 1.3, 330, DESIGN_W - 2.6 * MARGIN, 460, s.get("quote") or "", 58, ctx.fg,
              italic=True, font=ctx.display)
    if s.get("attribution"):
        _text_box(slide, ctx, MARGIN * 1.3, 820, DESIGN_W - 2.6 * MARGIN, 80, "— " + s["attribution"], 38, ctx.muted)


def _layout_big_fact(slide, s, ctx):
    _text_box(slide, ctx, MARGIN, 280, DESIGN_W - 2 * MARGIN, 260, s.get("fact") or "", 176, ctx.accent,
              bold=True, font=ctx.display, align="center")
    if s.get("label"):
        _text_box(slide, ctx, MARGIN, 580, DESIGN_W - 2 * MARGIN, 130, s["label"], 46, ctx.fg, align="center")
    if s.get("caption"):
        _text_box(slide, ctx, MARGIN, 730, DESIGN_W - 2 * MARGIN, 90, s["caption"], 30, ctx.muted, align="center")


def _layout_data(slide, s, ctx):
    _title(slide, ctx, s.get("title"))
    top = 290 if s.get("title") else MARGIN
    _chart(slide, ctx, s.get("chart") or {"empty": "No chart data"}, MARGIN, top, DESIGN_W - 2 * MARGIN,
           DESIGN_H - top - MARGIN - (70 if s.get("caption") else 0))
    if s.get("caption"):
        _text_box(slide, ctx, MARGIN, DESIGN_H - MARGIN - 50, DESIGN_W - 2 * MARGIN, 50, s["caption"], 28, ctx.muted)


def _layout_blank(slide, s, ctx):
    if s.get("title"):
        _title(slide, ctx, s["title"], size=64, x=MARGIN, y=440, w=DESIGN_W - 2 * MARGIN, h=200, align="center")


def _value_lines(lines):
    out = []
    for line in lines:
        for sep in (" — ", " – ", " - ", ": ", " | "):
            if sep in line:
                value, label = line.split(sep, 1)
                out.append((value.strip(), label.strip()))
                break
        else:
            out.append((line.strip(), ""))
    return out


def _layout_stats(slide, s, ctx):
    _title(slide, ctx, s.get("title"))
    items = _value_lines((s.get("body") or [])[:4])
    if not items:
        return
    col_w = (DESIGN_W - 2 * MARGIN) / len(items)
    for i, (value, label) in enumerate(items):
        x = MARGIN + i * col_w
        _text_box(slide, ctx, x, 420, col_w, 180, value, 124 if len(items) <= 3 else 104, ctx.accent,
                  bold=True, font=ctx.display, align="center")
        _text_box(slide, ctx, x + 20, 620, col_w - 40, 140, label, 34, ctx.fg, align="center")
        if i < len(items) - 1:
            _rect(slide, x + col_w - 1, 430, 2, 200, ctx.muted, alpha=30)


def _layout_steps(slide, s, ctx):
    _title(slide, ctx, s.get("title"))
    steps = (s.get("body") or [])[:6]
    if not steps:
        return
    if len(steps) <= 4:
        col_w = (DESIGN_W - 2 * MARGIN - 44 * (len(steps) - 1)) / len(steps)
        for i, step in enumerate(steps):
            x = MARGIN + i * (col_w + 44)
            badge = _oval(slide, x, 330, 76, ctx.accent)
            _fill_text(badge.text_frame, str(i + 1), ctx, 38, ctx.accent_contrast, bold=True, align="center")
            _rect(slide, x, 432, col_w, 3, ctx.accent, alpha=35)
            _text_box(slide, ctx, x, 460, col_w, 400, step, 34, ctx.fg)
    else:
        for i, step in enumerate(steps):
            y = 300 + i * 100
            badge = _oval(slide, MARGIN, y, 56, ctx.accent)
            _fill_text(badge.text_frame, str(i + 1), ctx, 28, ctx.accent_contrast, bold=True, align="center")
            _text_box(slide, ctx, MARGIN + 84, y + 6, DESIGN_W - 2 * MARGIN - 84, 70, step, 32, ctx.fg)


def _layout_timeline(slide, s, ctx):
    _title(slide, ctx, s.get("title"))
    events = _value_lines((s.get("body") or [])[:6])
    if not events:
        return
    col_w = (DESIGN_W - 2 * MARGIN) / len(events)
    line_y = 520
    _rect(slide, MARGIN, line_y - 2, DESIGN_W - 2 * MARGIN, 4, ctx.muted, alpha=35)
    for i, (date, label) in enumerate(events):
        cx = MARGIN + i * col_w + col_w / 2
        _text_box(slide, ctx, cx - col_w / 2, line_y - 100, col_w, 60, date, 40, ctx.accent, bold=True,
                  font=ctx.display, align="center")
        _oval(slide, cx - 14, line_y - 14, 28, ctx.accent)
        _text_box(slide, ctx, cx - col_w / 2 + 10, line_y + 40, col_w - 20, 200, label, 30, ctx.fg, align="center")


def _layout_image_grid(slide, s, ctx):
    _title(slide, ctx, s.get("title"))
    refs = (s.get("images") or [])[:4]
    if not refs:
        return
    top, gap = 290, 28
    area_w, area_h = DESIGN_W - 2 * MARGIN, DESIGN_H - top - MARGIN
    if len(refs) == 4:
        cell_w, cell_h = (area_w - gap) / 2, (area_h - gap) / 2
        for i, ref in enumerate(refs):
            r, c = divmod(i, 2)
            _picture(slide, ctx, ref, MARGIN + c * (cell_w + gap), top + r * (cell_h + gap), cell_w, cell_h)
    else:
        cell_w = (area_w - gap * (len(refs) - 1)) / len(refs)
        for i, ref in enumerate(refs):
            _picture(slide, ctx, ref, MARGIN + i * (cell_w + gap), top, cell_w, area_h)


LAYOUTS = {
    "title": _layout_title,
    "section": _layout_section,
    "bullets": _layout_bullets,
    "twoColumn": _layout_two_column,
    "imageLeft": lambda sl, s, c: _layout_image_side(sl, s, c, True),
    "imageRight": lambda sl, s, c: _layout_image_side(sl, s, c, False),
    "fullBleed": _layout_full_bleed,
    "quote": _layout_quote,
    "bigFact": _layout_big_fact,
    "chart": _layout_data,
    "table": _layout_data,
    "blank": _layout_blank,
    "stats": _layout_stats,
    "steps": _layout_steps,
    "timeline": _layout_timeline,
    "imageGrid": _layout_image_grid,
}

# Which python-pptx default-template layout carries a real title placeholder for each.
PLACEHOLDER_LAYOUT = {"title": 0}
TITLE_ONLY = {"section", "bullets", "twoColumn", "imageLeft", "imageRight", "chart", "table", "stats", "steps",
              "timeline", "imageGrid", "blank"}


def _background(slide, s, ctx):
    fill = slide.background.fill
    override = s.get("background")
    gradient = ctx.theme.get("gradient")
    if override is None and gradient:
        fill.gradient()
        fill.gradient_angle = float(gradient.get("angle", 90))
        stops = fill.gradient_stops
        stops[0].color.rgb = _rgb(gradient.get("from"))
        stops[1].color.rgb = _rgb(gradient.get("to"))
        return
    fill.solid()
    fill.fore_color.rgb = _rgb(ctx.accent if override == "accent" else ctx.surface if override == "surface"
                               else ctx.background)


def _footer(slide, s, ctx, number: Optional[int]):
    quiet = s.get("layout") in ("title", "section", "fullBleed")
    logo = ctx.picture(ctx.brand.get("logo"))
    x = MARGIN * 0.6
    if logo is not None:
        from PIL import Image
        with Image.open(logo) as im:
            w, h = im.size
        height = 48.0
        width = min(220.0, w * height / max(h, 1))
        slide.shapes.add_picture(str(logo), _emu(x), _emu(DESIGN_H - 36 - height), _emu(width), _emu(height))
        x += width + 20
    if quiet:
        return
    if ctx.brand.get("footerText"):
        _text_box(slide, ctx, x, DESIGN_H - 36 - 34, 900, 34, ctx.brand["footerText"], 22, ctx.muted)
    if number is not None:
        _text_box(slide, ctx, DESIGN_W - MARGIN * 0.6 - 100, DESIGN_H - 36 - 34, 100, 34, str(number), 22,
                  ctx.muted, align="right")


def _register_notes_master(presentation):
    """python-pptx adds a notes master the first time a slide gets notes, but never lists it in
    presentation.xml's p:notesMasterIdLst. PowerPoint tolerates that; Apple's Office importer
    (Quick Look, Finder previews) refuses to render the file at all. Register it."""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn
    root = presentation.part._element
    if root.find(qn("p:notesMasterIdLst")) is not None:
        return
    r_id = next((rid for rid, rel in presentation.part.rels.items() if rel.reltype == RT.NOTES_MASTER), None)
    if r_id is None:
        return
    id_list = etree.Element(qn("p:notesMasterIdLst"))
    entry = etree.SubElement(id_list, qn("p:notesMasterId"))
    entry.set(qn("r:id"), r_id)
    root.find(qn("p:sldMasterIdLst")).addnext(id_list)


def build(deck: dict, assets_dir: Optional[str]):
    _pptx()
    from pptx import Presentation
    from pptx.util import Inches

    ctx = Ctx(deck, assets_dir)
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W_IN)
    presentation.slide_height = Inches(SLIDE_H_IN)
    presentation.core_properties.title = deck.get("title") or "Slides"

    slides = deck.get("slides") or []
    for s in slides:
        layout_id = s.get("layout") or "bullets"
        if layout_id in PLACEHOLDER_LAYOUT:
            layout = presentation.slide_layouts[PLACEHOLDER_LAYOUT[layout_id]]
        elif layout_id in TITLE_ONLY and s.get("title"):
            layout = presentation.slide_layouts[5]
        else:
            layout = presentation.slide_layouts[6]
        slide = presentation.slides.add_slide(layout)
        # The default template's other placeholders (the title slide's subtitle) would show
        # "Click to add subtitle" in PowerPoint; this module draws its own.
        for placeholder in list(slide.placeholders):
            if placeholder.placeholder_format.idx != 0:
                placeholder._element.getparent().remove(placeholder._element)
        _background(slide, s, ctx)
        LAYOUTS.get(layout_id, _layout_bullets)(slide, s, ctx)
        if deck.get("brand"):
            _footer(slide, s, ctx, s.get("number"))
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    _register_notes_master(presentation)
    return presentation, ctx.warnings, len(slides)


# --------------------------------------------------------------------------- routes


@router.get("/healthz")
def healthz():
    return {"status": "ok"}


@router.get("/info")
def info():
    try:
        import pptx
        version = getattr(pptx, "__version__", "unknown")
    except ImportError:
        version = None
    return {"service": "slides-pptx", "python_pptx": version, "layouts": sorted(LAYOUTS)}


@router.post("/prepare")
def prepare():
    """Makes sure python-pptx is importable (installing it if needed) — the app calls this
    right after installing the service so the first export does not pay for it."""
    pptx = _pptx()
    return {"status": "ok", "python_pptx": getattr(pptx, "__version__", "unknown")}


@router.post("/export")
def export(request: ExportRequest):
    output = Path(request.output_path).expanduser()
    if output.suffix.lower() != ".pptx":
        raise HTTPException(status_code=400, detail="output_path must end in .pptx")
    if not output.parent.exists():
        raise HTTPException(status_code=400, detail=f"folder does not exist: {output.parent}")
    presentation, warnings, count = build(request.deck, request.assets_dir)
    partial = output.with_name(output.name + ".part")
    presentation.save(str(partial))
    os.replace(partial, output)   # rename only on a COMPLETE write
    return {"status": "ok", "output_path": str(output), "slides": count,
            "size_bytes": output.stat().st_size, "warnings": sorted(set(warnings))}
